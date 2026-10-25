#!/usr/bin/env python3
"""One-shot (2026-10-25): replace the body of Chapter 19 (HOW HYPNOSIS REALLY
WORKS) in Built-for-Wonder.docx with Chris's approved rewrite from
drafts/ch19-how-hypnosis-really-works-2026-10-25.md.

Keeps: the chapter opener block (Heading1, hook line, tier badges) and the
"CHAPTER 19 — ..." label paragraph. Replaces: everything from the drop-cap
paragraph through the key-read line. Formatting is cloned from the chapter's
own paragraphs (drop cap, body, header, quote, Take, separator, key read).
Backup: backups/Built-for-Wonder.pre-ch19-rewrite-2026-10-25.docx
"""
import copy, re, shutil, sys
from pathlib import Path
from docx import Document

ROOT = Path(__file__).parent
DOCX = ROOT / "Built-for-Wonder.docx"
DRAFT = ROOT / "drafts" / "ch19-how-hypnosis-really-works-2026-10-25.md"
BACKUP = ROOT / "backups" / "Built-for-Wonder.pre-ch19-rewrite-2026-10-25.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

HEADERS = {
    "Why People Find Hypnosis Hard to Believe",
    "A Cold War Memo That Stayed With Me",
    "What Hypnosis Actually Is, and the Argument Behind It",
    "Hypnotizability Is the Real Thing",
    "Who Responds, and Why Absorption Is Not the Answer You Think",
    "The Induction Adds Less Than You Think",
    "The Wording Decides Which Part of the Brain Does the Work",
    "What the Brain Scans Show, and What They Do Not",
    "Why It Feels Like It Is Happening by Itself",
    "Where the Swinging Watch Comes From",
    "What This Means for a Stage Show",
}

def smart(text):
    """Straight quotes -> curly, per paragraph."""
    text = re.sub(r"(\w)'(\w)", "\u2019".join(["\\1", "\\2"]), text)   # it's -> it’s
    text = re.sub(r"(\w)'", "\\1\u2019", text)                            # trailing apostrophes
    out, open_q = [], True
    for ch in text:
        if ch == '"':
            out.append("\u201c" if open_q else "\u201d"); open_q = not open_q
        else:
            out.append(ch)
    return "".join(out)

def load_blocks():
    lines = DRAFT.read_text(encoding="utf-8").split("\n")
    start = next(i for i, l in enumerate(lines) if l.startswith("CHAPTER 19"))
    blocks = []
    for l in lines[start + 1:]:
        s = l.strip()
        if not s or s.startswith("[CHRIS:"):
            continue
        blocks.append(s)
    return blocks

def set_text(p_elem, text):
    runs = p_elem.findall(W + "r")
    for r in runs[1:]:
        p_elem.remove(r)
    ts = runs[0].findall(W + "t")
    for t in ts[1:]:
        runs[0].remove(t)
    ts[0].text = text
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return p_elem

def main():
    shutil.copy2(DOCX, BACKUP)
    print("backup ->", BACKUP.name)
    d = Document(str(DOCX))
    ps = d.paragraphs
    label_i = next(i for i, p in enumerate(ps) if p.text.startswith("CHAPTER 19") and "HYPNOSIS" in p.text)
    keyread_i = next(i for i in range(label_i, len(ps)) if ps[i].text.strip() == "Trance is just the brain doing what it already does every day.")
    assert ps[keyread_i + 1].text.strip() == "PART THREE", ps[keyread_i + 1].text
    # templates (clone before deleting)
    T = {}
    T["drop"] = copy.deepcopy(ps[label_i + 1]._p);   assert ps[label_i + 1].text == "W"
    T["droprest"] = copy.deepcopy(ps[label_i + 2]._p)
    T["feels"] = copy.deepcopy(ps[label_i + 3]._p);  assert ps[label_i + 3].text == "WHAT IT FEELS LIKE"
    T["body"] = copy.deepcopy(ps[label_i + 4]._p)
    hdr = next(p for p in ps[label_i:keyread_i] if p.text == "Why People Find Hypnosis Hard to Believe")
    T["header"] = copy.deepcopy(hdr._p)
    quote = next(p for p in ps[label_i:keyread_i] if p.text.startswith("\u201cFrankly, I now distrust"))
    T["quote"] = copy.deepcopy(quote._p)
    take = next(p for p in ps[label_i:keyread_i] if p.text == "Chris Michael's Take")
    T["take"] = copy.deepcopy(take._p)
    sep = ps[keyread_i - 1]; assert sep.text.strip() == "\u00b7 \u00b7 \u00b7", sep.text
    T["sep"] = copy.deepcopy(sep._p)
    T["keyread"] = copy.deepcopy(ps[keyread_i]._p)
    for k, e in T.items():
        assert len(e.findall(W + "r")) >= 1, k

    # delete old body
    old = ps[label_i + 1:keyread_i + 1]
    print(f"removing {len(old)} paragraphs (idx {label_i+1}..{keyread_i})")
    for p in old:
        p._p.getparent().remove(p._p)

    # build new
    blocks = load_blocks()
    new = []
    first_body_done = False
    in_take = False
    for i, s in enumerate(blocks):
        last = (i == len(blocks) - 1)
        if s == "WHAT IT FEELS LIKE":
            new.append(set_text(copy.deepcopy(T["feels"]), s)); continue
        if s == "Chris Michael's Take":
            in_take = True
            new.append(set_text(copy.deepcopy(T["take"]), s)); continue
        if s == "\u00b7 \u00b7 \u00b7":
            new.append(set_text(copy.deepcopy(T["sep"]), s)); continue
        if last:
            new.append(set_text(copy.deepcopy(T["keyread"]), s)); continue
        if s in HEADERS:
            new.append(set_text(copy.deepcopy(T["header"]), s)); continue
        if s.startswith('"Frankly, I now distrust'):
            new.append(set_text(copy.deepcopy(T["quote"]), smart(s))); continue
        if not first_body_done:
            first_body_done = True
            new.append(set_text(copy.deepcopy(T["drop"]), s[0]))
            new.append(set_text(copy.deepcopy(T["droprest"]), smart(s[1:])))
            continue
        new.append(set_text(copy.deepcopy(T["body"]), smart(s)))

    ref = ps[label_i]._p
    for e in reversed(new):
        ref.addnext(e)
    d.save(str(DOCX))
    print(f"inserted {len(new)} paragraphs")

    # verify
    d2 = Document(str(DOCX)); ps2 = d2.paragraphs
    li = next(i for i, p in enumerate(ps2) if p.text.startswith("CHAPTER 19") and "HYPNOSIS" in p.text)
    for p in ps2[li:li + len(new) + 3]:
        b = any(r.bold for r in p.runs)
        print(f"  {'B' if b else ' '} {p.text[:78]}")

if __name__ == "__main__":
    main()
